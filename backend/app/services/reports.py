"""Statutory compliance reports (FR13): aggregation + PDF (ReportLab) + Excel (openpyxl)."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime, time
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import Select, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import Forbidden, NotFound
from app.models import (
    AttendanceRecord,
    ComplianceItem,
    CorrectiveAction,
    Inspection,
    Mine,
    Report,
    Subsidiary,
    Violation,
)
from app.models.enums import ActionStatus, ComplianceStatus, ReportScope, Role, ViolationStatus
from app.schemas.governance import ReportCreate
from app.services.access import Perm, Principal, assert_mine_access
from app.services.events import record_event

GREEN = colors.HexColor("#0b3d2e")


def _label(key: str) -> str:
    return key.replace("_pct", " (%)").replace("_", " ").capitalize()


def report_workflow_id(report_id: uuid.UUID | str) -> str:
    return f"report-generation-{report_id}"


async def scope_label(session: AsyncSession, scope: ReportScope, scope_id: uuid.UUID | None) -> str:
    if scope == ReportScope.NATIONAL:
        return "All subsidiaries (National)"
    if scope == ReportScope.SUBSIDIARY:
        sub = await session.get(Subsidiary, scope_id)
        return f"{sub.name} ({sub.code})" if sub else "Unknown subsidiary"
    mine = await session.get(Mine, scope_id)
    return f"{mine.name} ({mine.code})" if mine else "Unknown mine"


async def _mine_ids(session: AsyncSession, scope: ReportScope, scope_id: uuid.UUID | None) -> list[uuid.UUID]:
    stmt = select(Mine.id)
    if scope == ReportScope.SUBSIDIARY:
        stmt = stmt.where(Mine.subsidiary_id == scope_id)
    elif scope == ReportScope.MINE:
        stmt = stmt.where(Mine.id == scope_id)
    return list((await session.execute(stmt)).scalars().all())


async def create_request(session: AsyncSession, principal: Principal, data: ReportCreate) -> Report:
    principal.require(Perm.REPORT_GENERATE)
    if data.scope == ReportScope.MINE:
        await assert_mine_access(session, principal, data.scope_id)  # type: ignore[arg-type]
    elif data.scope == ReportScope.SUBSIDIARY:
        if not principal.is_global and principal.subsidiary_id != data.scope_id:
            raise Forbidden("You can only generate reports for your own subsidiary")
    elif not principal.is_global:
        raise Forbidden("National reports require corporate HQ, regulator or admin access")
    label = await scope_label(session, data.scope, data.scope_id)
    report = Report(
        title=f"Compliance Report — {label} — {data.period_start:%d %b %Y} to {data.period_end:%d %b %Y}",
        scope=data.scope, scope_id=data.scope_id, period_start=data.period_start, period_end=data.period_end,
        requested_by=principal.id,
    )
    session.add(report)
    await session.flush()
    report.workflow_id = report_workflow_id(report.id)
    record_event(session, "report.requested", entity_type="report", entity_id=report.id, actor=principal,
                 data={"scope": data.scope.value, "scope_id": str(data.scope_id) if data.scope_id else None})
    await session.commit()
    return report


def report_query(principal: Principal) -> Select[Any]:
    stmt = select(Report)
    if principal.is_global:
        return stmt
    if principal.role == Role.MINE_OFFICIAL:
        return stmt.where((Report.scope == ReportScope.MINE) & (Report.scope_id == principal.mine_id))
    if principal.role == Role.CORPORATE:
        mines = select(Mine.id).where(Mine.subsidiary_id == principal.subsidiary_id)
        return stmt.where(
            ((Report.scope == ReportScope.SUBSIDIARY) & (Report.scope_id == principal.subsidiary_id))
            | ((Report.scope == ReportScope.MINE) & Report.scope_id.in_(mines))
        )
    return stmt.where(Report.requested_by == principal.id)


async def get_report(session: AsyncSession, principal: Principal, report_id: uuid.UUID) -> Report:
    report = (await session.execute(report_query(principal).where(Report.id == report_id))).scalars().first()
    if report is None:
        raise NotFound("Report not found")
    return report


async def aggregate(session: AsyncSession, scope: ReportScope, scope_id: uuid.UUID | None,
                    start: date, end: date) -> dict[str, Any]:
    mine_ids = await _mine_ids(session, scope, scope_id)
    t0 = datetime.combine(start, time.min, tzinfo=UTC)
    t1 = datetime.combine(end, time.max, tzinfo=UTC)
    if not mine_ids:
        return {"mines": [], "totals": {}}
    comp_rows = (
        await session.execute(
            select(ComplianceItem.status, func.count()).where(ComplianceItem.mine_id.in_(mine_ids))
            .group_by(ComplianceItem.status)
        )
    ).all()
    comp = {s.value: 0 for s in ComplianceStatus}
    for st, n in comp_rows:
        comp[st.value] = int(n)
    v_rows = (
        await session.execute(
            select(Violation.severity, Violation.status, Violation.category, func.count())
            .where(Violation.mine_id.in_(mine_ids), Violation.occurred_at.between(t0, t1))
            .group_by(Violation.severity, Violation.status, Violation.category)
        )
    ).all()
    by_sev: dict[str, int] = {}
    by_cat: dict[str, int] = {}
    closed = 0
    total_v = 0
    for sev, st, cat, n in v_rows:
        by_sev[sev.value] = by_sev.get(sev.value, 0) + n
        by_cat[cat.value] = by_cat.get(cat.value, 0) + n
        total_v += n
        if st == ViolationStatus.CLOSED:
            closed += n
    insp = (
        await session.execute(
            select(Inspection.outcome, func.count()).where(Inspection.mine_id.in_(mine_ids),
                                                           Inspection.inspected_at.between(t0, t1))
            .group_by(Inspection.outcome)
        )
    ).all()
    insp_by = {o.value: int(n) for o, n in insp}
    geo_ok = (
        await session.execute(
            select(func.count()).where(Inspection.mine_id.in_(mine_ids), Inspection.inspected_at.between(t0, t1),
                                       Inspection.geo_verified.is_(True))
        )
    ).scalar_one()
    overdue_actions = (
        await session.execute(
            select(func.count()).select_from(CorrectiveAction).join(Violation)
            .where(Violation.mine_id.in_(mine_ids), CorrectiveAction.deadline < t1,
                   CorrectiveAction.status.not_in([ActionStatus.VERIFIED, ActionStatus.SUBMITTED]))
        )
    ).scalar_one()
    attendance = (
        await session.execute(
            select(func.count()).where(AttendanceRecord.mine_id.in_(mine_ids),
                                       AttendanceRecord.check_in_at.between(t0, t1))
        )
    ).scalar_one()
    per_mine_rows = (
        await session.execute(
            select(Mine.code, Mine.name, Mine.risk_score,
                   select(func.count()).where(Violation.mine_id == Mine.id, Violation.occurred_at.between(t0, t1))
                   .scalar_subquery(),
                   select(func.count()).where(Violation.mine_id == Mine.id, Violation.status != ViolationStatus.CLOSED)
                   .scalar_subquery(),
                   select(func.count()).where(ComplianceItem.mine_id == Mine.id,
                                              ComplianceItem.status == ComplianceStatus.OVERDUE).scalar_subquery(),
                   select(func.count()).where(Inspection.mine_id == Mine.id, Inspection.inspected_at.between(t0, t1))
                   .scalar_subquery())
            .where(Mine.id.in_(mine_ids)).order_by(Mine.risk_score.desc().nulls_last())
        )
    ).all()
    total_comp = sum(comp.values())
    return {
        "period": {"start": start.isoformat(), "end": end.isoformat()},
        "totals": {
            "mines": len(mine_ids),
            "compliance_items": total_comp,
            "compliance_rate_pct": round(100 * (comp["compliant"] + comp["due"]) / total_comp, 1) if total_comp else 100.0,
            "compliance_overdue": comp["overdue"],
            "compliance_violated": comp["violated"],
            "violations": total_v,
            "violations_closed": closed,
            "closure_rate_pct": round(100 * closed / total_v, 1) if total_v else 100.0,
            "inspections": sum(insp_by.values()),
            "geo_verified_inspections": int(geo_ok),
            "overdue_corrective_actions": int(overdue_actions),
            "attendance_records": int(attendance),
        },
        "compliance_by_status": comp,
        "violations_by_severity": by_sev,
        "violations_by_category": by_cat,
        "inspections_by_outcome": insp_by,
        "mines": [
            {"code": c, "name": n, "risk_score": r, "violations": v, "open_violations": ov,
             "overdue_compliance": oc, "inspections": i}
            for c, n, r, v, ov, oc, i in per_mine_rows
        ],
    }


def render_pdf(title: str, scope_label_: str, metrics: dict[str, Any], summary: str | None,
               audit_head: str | None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title, author="Lumen — Ministry of Coal")
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=GREEN, fontSize=16, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=GREEN, fontSize=12, spaceBefore=10)
    body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("s", parent=body, fontSize=8, textColor=colors.grey)
    story: list[Any] = [
        Paragraph("Government of India · Ministry of Coal · Lumen", small),
        Paragraph(title, h1),
        Paragraph(f"Scope: {scope_label_} &nbsp;|&nbsp; Generated: {datetime.now(UTC):%d %b %Y %H:%M} UTC", small),
        Spacer(1, 8),
    ]
    if summary:
        story += [Paragraph("Executive summary", h2)]
        story += [Paragraph(p.replace("&", "&amp;").replace("<", "&lt;"), body) for p in summary.split("\n") if p.strip()]

    def table(rows: list[list[Any]], widths: list[float] | None = None) -> Table:
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), GREEN),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f2")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cfd8d3")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return t

    totals = metrics.get("totals", {})
    story += [Paragraph("Key indicators", h2),
              table([["Indicator", "Value"]] + [[_label(k), str(v)] for k, v in totals.items()],
                    [110 * mm, 60 * mm])]
    for key, label in (("violations_by_severity", "Violations by severity"),
                       ("violations_by_category", "Violations by category"),
                       ("compliance_by_status", "Statutory compliance by status"),
                       ("inspections_by_outcome", "Inspections by outcome")):
        data = metrics.get(key) or {}
        if data:
            story += [Paragraph(label, h2),
                      table([["Class", "Count"]] + [[_label(k), str(v)] for k, v in data.items()],
                            [110 * mm, 60 * mm])]
    mines = metrics.get("mines", [])
    if mines:
        story += [Paragraph("Mine-wise status (ordered by risk)", h2),
                  table([["Code", "Mine", "Risk", "Violations", "Open", "Overdue", "Inspections"]] + [
                      [m["code"], m["name"][:34], f"{m['risk_score']:.0f}" if m["risk_score"] is not None else "—",
                       m["violations"], m["open_violations"], m["overdue_compliance"], m["inspections"]]
                      for m in mines
                  ])]
    story += [Spacer(1, 12), Paragraph(
        f"Integrity: audit-trail head hash at generation time {audit_head or 'n/a'}. "
        "This report is system-generated from the Lumen system of record.", small)]
    doc.build(story)
    return buf.getvalue()


def render_xlsx(title: str, metrics: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    header_fill = PatternFill("solid", fgColor="0B3D2E")
    header_font = Font(color="FFFFFF", bold=True)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Indicator", "Value"])
    for cell in ws[3]:
        cell.fill, cell.font = header_fill, header_font
    for k, v in metrics.get("totals", {}).items():
        ws.append([_label(k), v])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    def sheet(name: str, headers: list[str], rows: list[list[Any]]) -> None:
        s = wb.create_sheet(name)
        s.append(headers)
        for cell in s[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center")
        for r in rows:
            s.append(r)
        for i, _ in enumerate(headers, start=1):
            s.column_dimensions[get_column_letter(i)].width = 22
        s.freeze_panes = "A2"

    sheet("Mines", ["Code", "Mine", "Risk score", "Violations", "Open violations", "Overdue compliance", "Inspections"],
          [[m["code"], m["name"], m["risk_score"], m["violations"], m["open_violations"], m["overdue_compliance"],
            m["inspections"]] for m in metrics.get("mines", [])])
    sheet("Violations", ["Dimension", "Class", "Count"],
          [["severity", k, v] for k, v in metrics.get("violations_by_severity", {}).items()]
          + [["category", k, v] for k, v in metrics.get("violations_by_category", {}).items()])
    sheet("Compliance", ["Status", "Count"], [[k, v] for k, v in metrics.get("compliance_by_status", {}).items()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
